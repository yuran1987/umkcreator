from docx import Document
from .models import Competence
import re

############################################################################################
dump_test = False


class CompetenceImport(object):

    def __init__(self, fname):
        self.comp_list = []
        print("file: " + fname)

        document = Document(fname)
        wtable_prof_deyat = None
        wtable_lev_copeten = None

        # find table
        for table in document.tables:
            for row in table.rows:
                for cell in row.cells:
                    if str(cell.text) == u'Вид профессиональной деятельности':
                        wtable_prof_deyat = table
                    elif str(cell.text) == u'Уровень сформированности компетенции' or str(cell.text) == u'Уровни сформированности компетенции':
                        wtable_lev_copeten = table
                break

        if (wtable_prof_deyat is None) or (wtable_lev_copeten is None):
            print("---> tables not found")
            return None

        print('Starting analyse prof table...')
        self.table_data = dict()
        prof_record_name = None  # this is 1 column
        prof_record = None  # this is 1 column

        # skip first 3 rows
        rows_iter = iter(wtable_prof_deyat.rows)
        for i in range(3):
            next(rows_iter)

        for row in rows_iter:
            curr_prof_record_name = row.cells[0].text
            assert None != curr_prof_record_name
            if curr_prof_record_name != prof_record_name:
                prof_record_name = curr_prof_record_name
                assert not prof_record_name in self.table_data.keys()
                prof_record = dict()
                self.table_data[prof_record_name] = prof_record

            comptn_numb_name = row.cells[1].text
            self.comp_list.append(comptn_numb_name)
            assert not comptn_numb_name in prof_record.keys()
            prof_record[comptn_numb_name] = self.prepare_comp_info(row.cells)

        print('Starting analyse indicators table...')
        self.table_indicators = dict()
        table = []

        for row in range(1, len(wtable_lev_copeten.rows)):
            table.append({'0': wtable_lev_copeten.cell(row, 0).text, '1': wtable_lev_copeten.cell(row, 1).text,
                          '2': wtable_lev_copeten.cell(row, 2).text})

        for name in self.comp_list:
            indicators = self.get_indicators(table, name)
            self.table_indicators[name] = indicators

        print('Finish.')

    def get_indicators(self, table, competence):
        num_copt = '0'
        flag = False
        indicators_know = ""
        indicators_can = ""
        indicators_own = ""

        for row in table:
            s = row['0']
            if re.search(competence, s):  # competence - искомая компетенция
                res = re.findall(r"\d+", s)  # .split("-")[1]
                if (num_copt != res):
                    num_copt = res
                    flag = True
                else:
                    flag = False
            elif flag == True:
                tmp = re.findall(r"\d+", s)
                if tmp:
                    if tmp != num_copt:  # re.findall(r"\d+",s.split("-")[1])
                        break
                else:
                    if re.search(r'Знает', row['1']):
                        indicators_know += "{0}; ".format(row['2'])
                    elif re.search(r'Умеет', row['1']):
                        indicators_can += "{0}; ".format(row['2'])
                    elif re.search(r'Владеет', row['1']):
                        indicators_own += "{0}; ".format(row['2'])

        return {'indicators_know': indicators_know, 'indicators_can': indicators_can, 'indicators_own': indicators_own}

    def prepare_comp_info(self, cells):
        result_data = dict()

        result_data["description"] = cells[2].text.strip()
        result_data["significance"] = cells[3].text.strip()

        result_data["should_know"] = cells[4].text.strip()
        result_data["should_able"] = cells[5].text.strip()
        result_data["should_master"] = cells[6].text.strip()

        dat = cells[7].text
        result_data["disciplines"] = [x.strip() for x in list(filter(None, dat.split(',')))]
        result_data["methods"] = cells[8].text.strip()
        return result_data

    def get_discipline_list(self):
        disciplines = set()
        for prof, prof_record in self.table_data.items():
            for comp, comp_record in prof_record.items():
                for discipline in comp_record['disciplines']:
                    disciplines.add(discipline.lower())
        return disciplines

    def find_competence(self, competence):
        for prof, prof_record in self.table_data.items():
            for comp, comp_record in prof_record.items():
                if competence == comp:
                    return comp_record
        return None


    #---------------print functions----------------------------------------
    def get_competence_description(self, competence): #возвращает поля знать уметь владеть для заданной компетенции
        comp_record = self.find_competence(competence)
        if None == comp_record:
            print('Competence "' + competence + '" not found')
            return

        return {'should_know': comp_record['should_know'],       #Должен знать:
               'should_able': comp_record['should_able'],       #Должен уметь:
               'should_master': comp_record['should_master']    #Должен владеть:
               }

    def get_methods(self, competences = []):
        competences = list(set(competences))

        methods = set()
        for competence in competences:
            comp_descr = self.find_competence(competence)
            assert None != comp_descr
            methods.add(comp_descr['methods'])

        return methods

    def get_description(self, competence):
        comp_record = self.find_competence(competence)
        if None == comp_record:
            print('Competence "' + competence + '" not found')
            return

        return comp_record['description']



    def print_find_discipline(self, discipline):
        disciplines = self.get_discipline_list()
        if discipline.lower() in disciplines:
            print('Дисциплина: "' + discipline + '" найдена')
        else:
            print('Дисциплина: "' + discipline + '" ненайдена')
###############################################################################################################
    def run_import(self, direc, tr_prog, qualif):

        for name in self.comp_list:
            comp_record = self.find_competence(name)
            indicators = self.table_indicators[name]
            #print(indicators['indicators_know'])

            comp_new = Competence(name=name, direction=direc, full_content=comp_record['description'],
                                  should_know=comp_record['should_know'],
                                  should_able=comp_record['should_able'],
                                  should_master=comp_record['should_master'],
                                  training_program=tr_prog, qualif=qualif,
                                  indicators_know=indicators['indicators_know'],
                                  indicators_can=indicators['indicators_can'],
                                  indicators_own=indicators['indicators_own'])
            comp_new.save()

